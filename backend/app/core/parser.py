import io
import re
from typing import Tuple, Dict, Any, Optional, Callable
import docx
import pypdf
from app.core.visual_audit import analyze_image_asset
FILE_SIGNATURES = {'pdf': [b'%PDF-'], 'docx': [b'PK\x03\x04'], 'png': [b'\x89PNG\r\n\x1a\n'], 'jpg': [b'\xff\xd8\xff'], 'jpeg': [b'\xff\xd8\xff']}
DANGEROUS_PREAMBLES = (b'MZ', b'\x7fELF', b'#!')

class SignatureVerifier:

    @classmethod
    def verify(cls, target_path: str, byte_payload: bytes) -> str:
        if not byte_payload:
            raise ValueError('Uploaded file is empty (0 bytes).')
        segments = target_path.lower().rsplit('.', 1)
        if len(segments) < 2 or not segments[1]:
            raise ValueError('Missing file extension.')
        ext_label = segments[1]
        if any((byte_payload.startswith(p) for p in DANGEROUS_PREAMBLES)):
            raise ValueError('Dangerous file signature detected: binary executables and scripts are prohibited.')
        if ext_label in FILE_SIGNATURES:
            allowed_prefixes = FILE_SIGNATURES[ext_label]
            if not any((byte_payload.startswith(pfx) for pfx in allowed_prefixes)):
                raise ValueError(f"File signature mismatch for '.{ext_label}': deceptive or corrupted file format.")
        elif ext_label == 'txt':
            if b'\x00' in byte_payload[:4096]:
                raise ValueError('Binary content detected in plain text file.')
        else:
            raise ValueError(f"Unsupported file extension '.{ext_label}'. Supported: .pdf, .docx, .txt, .png, .jpg, .jpeg")
        return ext_label

def validate_file_signature(filename: str, content_bytes: bytes) -> str:
    return SignatureVerifier.verify(filename, content_bytes)

class TextExtractionPipeline:

    @staticmethod
    def decode_plain(byte_payload: bytes) -> str:
        try:
            return byte_payload.decode('utf-8')
        except UnicodeDecodeError:
            return byte_payload.decode('latin-1', errors='replace')

def extract_text_from_txt(content_bytes: bytes) -> str:
    return TextExtractionPipeline.decode_plain(content_bytes)

class DocxTextExtractor:

    @staticmethod
    def extract_prose(byte_payload: bytes) -> str:
        try:
            document_obj = docx.Document(io.BytesIO(byte_payload))
            lines = [p.text for p in document_obj.paragraphs if p.text.strip()]
            for tbl in document_obj.tables:
                for row_node in tbl.rows:
                    cell_segments = [c.text.strip() for c in row_node.cells if c.text.strip()]
                    if cell_segments:
                        lines.append(' | '.join(cell_segments))
            return '\n\n'.join(lines)
        except Exception as err:
            raise ValueError(f'Malformed or corrupted DOCX document: {str(err)}')

def extract_text_from_docx(content_bytes: bytes) -> str:
    return DocxTextExtractor.extract_prose(content_bytes)

class PdfTextExtractor:

    @staticmethod
    def extract_pages(byte_payload: bytes) -> str:
        try:
            pdf_engine = pypdf.PdfReader(io.BytesIO(byte_payload))
            aggregated = []
            for sheet in pdf_engine.pages:
                try:
                    slice_content = sheet.extract_text()
                    if slice_content:
                        aggregated.append(slice_content.strip())
                except Exception:
                    continue
            if not aggregated:
                return '[PDF scanned or empty text layer. Visual inspection required.]'
            return '\n\n'.join(aggregated)
        except Exception as fault:
            raise ValueError(f'Malformed or corrupted PDF document: {str(fault)}')

def extract_text_from_pdf(content_bytes: bytes) -> str:
    return PdfTextExtractor.extract_pages(content_bytes)

class UploadedFileParser:

    @staticmethod
    def _render_image_summary(content: bytes, tag: str) -> str:
        metrics = analyze_image_asset(content, tag)
        swatch_sample = ', '.join((entry['hex'] for entry in metrics['dominant_colors'][:3]))
        return f"[Image Asset: {tag} ({metrics['width']}x{metrics['height']}px, format: {metrics['format']}). Aspect ratio: {metrics['aspect_ratio']}:1. Dominant palette: {swatch_sample}. Visual brand and contrast checks completed.]"
    PARSER_DISPATCH: Dict[str, Callable[[bytes, str], str]] = {'txt': lambda b, _: extract_text_from_txt(b), 'docx': lambda b, _: extract_text_from_docx(b), 'pdf': lambda b, _: extract_text_from_pdf(b), 'png': lambda b, tag: UploadedFileParser._render_image_summary(b, tag), 'jpg': lambda b, tag: UploadedFileParser._render_image_summary(b, tag), 'jpeg': lambda b, tag: UploadedFileParser._render_image_summary(b, tag)}

    @classmethod
    def parse(cls, target_name: str, byte_payload: bytes) -> str:
        ext_verified = SignatureVerifier.verify(target_name, byte_payload)
        handler = cls.PARSER_DISPATCH.get(ext_verified)
        if not handler:
            raise ValueError(f"Unsupported file type '.{ext_verified}'. Supported: .pdf, .docx, .txt, .png, .jpg, .jpeg")
        return handler(byte_payload, target_name)

def parse_uploaded_file(filename: str, content_bytes: bytes) -> str:
    return UploadedFileParser.parse(filename, content_bytes)

class BoundedStreamReader:
    BLOCK_CAPACITY = 64 * 1024

    @classmethod
    async def read_bounded(cls, source_stream: Any, size_ceiling_mb: int) -> bytes:
        max_bytes_threshold = size_ceiling_mb * 1024 * 1024
        stream_chunks = []
        bytes_ingested = 0
        while True:
            slice_data = await source_stream.read(cls.BLOCK_CAPACITY)
            if not slice_data:
                break
            bytes_ingested += len(slice_data)
            if bytes_ingested > max_bytes_threshold:
                raise ValueError(f'Uploaded file exceeds maximum permitted size of {size_ceiling_mb}MB.')
            stream_chunks.append(slice_data)
        assembled = b''.join(stream_chunks)
        if not assembled:
            raise ValueError('Uploaded file is empty (0 bytes).')
        return assembled

async def read_uploaded_file_safely(file: Any, max_size_mb: int) -> bytes:
    return await BoundedStreamReader.read_bounded(file, max_size_mb)
